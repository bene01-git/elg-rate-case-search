#!/usr/bin/env python3
from __future__ import annotations
"""
DPS Document PDF -> Searchable PDF + Text Extractor
====================================================
Reads the ELG_Database.xlsx file, downloads PDFs from the 'Link to Document'
column, converts them to text-searchable PDFs using OCR, and extracts the
full text into .txt files.

Output structure:
    searchable_pdfs/
        raw_downloads/      ← original PDFs as downloaded
        searchable/         ← OCR'd searchable PDFs
        text/               ← extracted text as .txt files

Usage:
    python pdf_converter.py [--input ELG_Database.xlsx] [--output-dir searchable_pdfs]
    python pdf_converter.py --case 25-G-0073
    python pdf_converter.py --download-only
    python pdf_converter.py --text-only       (skip OCR, just extract text from existing downloads)

Requirements:
    pip install openpyxl requests pymupdf
    For OCR: pip install ocrmypdf  (+ Tesseract installed on your system)
"""

import argparse
import os
import re
import sys
import time
import requests
from pathlib import Path

from openpyxl import load_workbook

try:
    import ocrmypdf
    HAS_OCR = True
except ImportError:
    HAS_OCR = False

try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

# Fallback text extractor
if not HAS_PYMUPDF:
    try:
        from PyPDF2 import PdfReader
        HAS_PYPDF2 = True
    except ImportError:
        HAS_PYPDF2 = False
else:
    HAS_PYPDF2 = False


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
DEFAULT_INPUT = "ELG_Database.xlsx"
DEFAULT_OUTPUT_DIR = "searchable_pdfs"
BASE_URL = "https://documents.dps.ny.gov"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
    )
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def sanitize_filename(name: str, max_len: int = 80) -> str:
    name = re.sub(r'[<>:"/\\|?*]', '_', name)
    name = re.sub(r'\s+', ' ', name).strip()
    if len(name) > max_len:
        name = name[:max_len].rstrip()
    return name


def download_pdf(url: str, dest_path: str, session: requests.Session) -> bool:
    try:
        resp = session.get(url, headers=HEADERS, timeout=60, stream=True, allow_redirects=True)
        resp.raise_for_status()

        content_type = resp.headers.get("Content-Type", "")
        if "html" in content_type.lower() and b"%PDF" not in resp.content[:10]:
            print(f"        [SKIP] URL returned HTML, not a PDF.")
            return False

        with open(dest_path, "wb") as f:
            for chunk in resp.iter_content(chunk_size=8192):
                f.write(chunk)

        with open(dest_path, "rb") as f:
            header = f.read(5)
        if header != b"%PDF-":
            print(f"        [SKIP] Downloaded file is not a valid PDF.")
            os.remove(dest_path)
            return False

        return True

    except requests.RequestException as e:
        print(f"        [ERROR] Download failed: {e}")
        return False


def ocr_pdf(input_path: str, output_path: str) -> bool:
    try:
        ocrmypdf.ocr(
            input_path,
            output_path,
            skip_text=True,
            deskew=True,
            optimize=1,
            progress_bar=False,
        )
        return True
    except ocrmypdf.exceptions.PriorOcrFoundError:
        import shutil
        shutil.copy2(input_path, output_path)
        print("already searchable.", end=" ")
        return True
    except ocrmypdf.exceptions.EncryptedPdfError:
        print("[SKIP] encrypted.", end=" ")
        return False
    except Exception as e:
        print(f"[ERROR] OCR failed: {e}", end=" ")
        return False


def extract_text_pymupdf(pdf_path: str) -> str:
    """Extract text from a PDF using PyMuPDF (best quality)."""
    doc = fitz.open(pdf_path)
    text_parts = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        page_text = page.get_text("text")
        if page_text.strip():
            text_parts.append(f"--- Page {page_num + 1} ---\n{page_text}")
    doc.close()
    return "\n\n".join(text_parts)


def extract_text_pypdf2(pdf_path: str) -> str:
    """Extract text from a PDF using PyPDF2 (fallback)."""
    reader = PdfReader(pdf_path)
    text_parts = []
    for i, page in enumerate(reader.pages):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
    return "\n\n".join(text_parts)


def extract_text(pdf_path: str) -> str:
    """Extract text from a PDF using the best available library."""
    if HAS_PYMUPDF:
        return extract_text_pymupdf(pdf_path)
    elif HAS_PYPDF2:
        return extract_text_pypdf2(pdf_path)
    else:
        return ""


def save_text(text: str, txt_path: str) -> bool:
    """Save extracted text to a .txt file. Returns True if text was found."""
    if not text.strip():
        return False
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(text)
    return True


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Download, OCR, and extract text from DPS case PDFs.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  python pdf_converter.py
  python pdf_converter.py --case 25-G-0073
  python pdf_converter.py --download-only
  python pdf_converter.py --text-only
  python pdf_converter.py --input ELG_Database.xlsx --output-dir searchable_pdfs
        """,
    )
    parser.add_argument("--input", "-i", default=DEFAULT_INPUT, help="Path to Excel file")
    parser.add_argument("--output-dir", "-o", default=DEFAULT_OUTPUT_DIR, help="Output directory")
    parser.add_argument("--download-only", action="store_true", help="Download PDFs without OCR or text extraction")
    parser.add_argument("--text-only", action="store_true", help="Skip OCR, just extract text from existing PDFs")
    parser.add_argument("--skip-ocr", action="store_true", help="Download and extract text, but skip OCR step")
    parser.add_argument("--case", "-c", default=None, help="Only process a specific case number")
    args = parser.parse_args()

    # Check dependencies
    if not HAS_PYMUPDF and not HAS_PYPDF2:
        print("[ERROR] No PDF text extraction library found.")
        print("        Install one with: pip install pymupdf")
        print("        (or as fallback: pip install PyPDF2)")
        if args.download_only:
            print("        Continuing in download-only mode...")
        else:
            sys.exit(1)

    if not HAS_OCR and not args.download_only and not args.text_only and not args.skip_ocr:
        print("[NOTE] ocrmypdf is not installed — OCR will be skipped.")
        print("       To enable OCR: pip install ocrmypdf (+ install Tesseract)")
        print("       Continuing with download + text extraction only...")
        print()
        args.skip_ocr = True

    # Read Excel
    if not os.path.exists(args.input):
        print(f"[ERROR] Excel file not found: {args.input}")
        sys.exit(1)

    wb = load_workbook(args.input)
    ws = wb.active

    rows_to_process = []
    for row_idx in range(2, ws.max_row + 1):
        case_number = ws.cell(row=row_idx, column=1).value or ""
        doc_type = ws.cell(row=row_idx, column=5).value or ""
        doc_title = ws.cell(row=row_idx, column=6).value or ""
        file_name = ws.cell(row=row_idx, column=8).value or ""

        cell = ws.cell(row=row_idx, column=10)
        link = ""
        if cell.hyperlink:
            link = cell.hyperlink.target or ""
        if not link:
            link = str(cell.value or "")
        link = link.strip()

        if not link or not link.startswith("http"):
            continue

        if args.case and args.case.lower() != case_number.lower():
            continue

        company_name = ws.cell(row=row_idx, column=2).value or ""    
        rows_to_process.append({
            "row": row_idx,
            "case_number": case_number,
            "company_name": company_name,
            "doc_type": doc_type,
            "doc_title": doc_title,
            "file_name": file_name,
            "link": link,
        })

    if not rows_to_process:
        print("[NOTE] No document links found in the spreadsheet.")
        if args.case:
            print(f"       (filtered to case: {args.case})")
        sys.exit(0)

    # Determine mode
    do_download = not args.text_only
    do_ocr = not args.download_only and not args.text_only and not args.skip_ocr and HAS_OCR
    do_text = not args.download_only and (HAS_PYMUPDF or HAS_PYPDF2)

    mode_parts = []
    if do_download:
        mode_parts.append("Download")
    if do_ocr:
        mode_parts.append("OCR")
    if do_text:
        mode_parts.append("Text extraction")
    mode_str = " + ".join(mode_parts) if mode_parts else "Download only"

    print("=" * 60)
    print("  DPS Document PDF Converter")
    print("=" * 60)
    print(f"  Input:      {args.input}")
    print(f"  Output dir: {args.output_dir}")
    print(f"  Documents:  {len(rows_to_process)}")
    if args.case:
        print(f"  Case:       {args.case}")
    print(f"  Mode:       {mode_str}")
    if HAS_PYMUPDF:
        print(f"  Text lib:   PyMuPDF")
    elif HAS_PYPDF2:
        print(f"  Text lib:   PyPDF2")
    print()

    # Create output directories
    raw_dir = os.path.join(args.output_dir, "raw_downloads")
    ocr_dir = os.path.join(args.output_dir, "searchable")
    txt_dir = os.path.join(args.output_dir, "text")

    if do_download:
        os.makedirs(raw_dir, exist_ok=True)
    if do_ocr:
        os.makedirs(ocr_dir, exist_ok=True)
    if do_text:
        os.makedirs(txt_dir, exist_ok=True)

    # Process each document
    session = requests.Session()
    stats = {"downloaded": 0, "ocr": 0, "text": 0, "skipped": 0, "failed": 0}

    for i, doc in enumerate(rows_to_process, 1):
        case_clean = sanitize_filename(doc["case_number"], max_len=20)
        company_clean = sanitize_filename(doc["company_name"], max_len=50)
        title_clean = sanitize_filename(doc["doc_title"] or doc["file_name"] or f"doc_{doc['row']}", max_len=60)
        type_clean = sanitize_filename(doc["doc_type"], max_len=40)

        base_name = f"{case_clean} - {company_clean} - {type_clean} - {title_clean}"
        if len(base_name) > 180:
            base_name = base_name[:180].rstrip()
        pdf_name = f"{base_name}.pdf"
        txt_name = f"{base_name}.txt"

        raw_path = os.path.join(raw_dir, pdf_name)
        ocr_path = os.path.join(ocr_dir, pdf_name)
        txt_path = os.path.join(txt_dir, txt_name)

        print(f"[{i}/{len(rows_to_process)}] {doc['doc_title'][:65]}")

        # --- Download ---
        if do_download:
            if os.path.exists(raw_path):
                print("        Downloaded: already exists.")
            else:
                print(f"        Downloading...", end=" ")
                ok = download_pdf(doc["link"], raw_path, session)
                if not ok:
                    stats["failed"] += 1
                    continue
                size_kb = os.path.getsize(raw_path) / 1024
                print(f"OK ({size_kb:.0f} KB)")
                stats["downloaded"] += 1

        # Determine which PDF to extract text from (prefer OCR'd version)
        source_pdf = raw_path

        # --- OCR ---
        if do_ocr:
            if os.path.exists(ocr_path):
                print("        OCR: already done.")
                source_pdf = ocr_path
            elif os.path.exists(raw_path):
                print("        OCR...", end=" ")
                ok = ocr_pdf(raw_path, ocr_path)
                if ok:
                    size_kb = os.path.getsize(ocr_path) / 1024
                    print(f"OK ({size_kb:.0f} KB)")
                    stats["ocr"] += 1
                    source_pdf = ocr_path
                else:
                    print()
        elif do_text and os.path.exists(ocr_path):
            # If text-only mode but OCR'd files exist from a previous run, use them
            source_pdf = ocr_path

        # --- Text extraction ---
        if do_text:
            if os.path.exists(txt_path):
                print("        Text: already extracted.")
                stats["skipped"] += 1
            elif os.path.exists(source_pdf):
                print("        Extracting text...", end=" ")
                text = extract_text(source_pdf)
                if save_text(text, txt_path):
                    word_count = len(text.split())
                    size_kb = os.path.getsize(txt_path) / 1024
                    print(f"OK ({word_count:,} words, {size_kb:.0f} KB)")
                    stats["text"] += 1
                else:
                    print("no text found (scanned image without OCR?).")
                    # If we skipped OCR and got no text, note it
                    if not do_ocr:
                        print("        Tip: run with OCR enabled to extract text from scanned pages.")
            else:
                print("        Text: no PDF available to extract from.")

        # Be polite to the server
        if do_download:
            time.sleep(1)

    # Summary
    print()
    print("=" * 60)
    print("  Summary")
    print("=" * 60)
    if do_download:
        print(f"  Downloaded:       {stats['downloaded']}")
    if do_ocr:
        print(f"  OCR'd:            {stats['ocr']}")
    if do_text:
        print(f"  Text extracted:   {stats['text']}")
    print(f"  Skipped (exists): {stats['skipped']}")
    print(f"  Failed:           {stats['failed']}")
    print()
    if do_download:
        print(f"  Raw PDFs:         {raw_dir}")
    if do_ocr:
        print(f"  Searchable PDFs:  {ocr_dir}")
    if do_text:
        print(f"  Text files:       {txt_dir}")


if __name__ == "__main__":
    main()
