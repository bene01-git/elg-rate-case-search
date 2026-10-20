#!/usr/bin/env python3
from __future__ import annotations
"""
NY DPS Rate Case Document Scraper
==================================
Scrapes documents from the NY Department of Public Service (DPS) Document
and Matter Management (DMM) system and appends them to an Excel workbook.

Accepts one or more case numbers/URLs in a single run.

Only captures the following document types:
  - Briefs
  - Joint Proposals and Stipulations
  - Orders
  - Petitions
  - Recommended Decisions
  - Testimony
  - Transcripts

Usage:
    python dps_scraper.py <case1> [case2] [case3] ... [--output path/to/ELG_Database.xlsx]

Examples:
    python dps_scraper.py 25-G-0073
    python dps_scraper.py 25-G-0073 23-G-0225 24-E-0165
    python dps_scraper.py 25-G-0073 "https://documents.dps.ny.gov/public/MatterManagement/CaseMaster.aspx?MatterSeq=75508&MNO=25-G-0073"
    python dps_scraper.py --file cases.txt
    python dps_scraper.py 25-G-0073 23-G-0225 --output ./ELG_Database.xlsx --visible

Requirements:
    pip install selenium openpyxl
"""

import argparse
import os
import re
import sys
import time

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    NoSuchElementException,
    StaleElementReferenceException,
)

from openpyxl import load_workbook
from openpyxl.styles import Font


# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
BASE_URL = "https://documents.dps.ny.gov"
CASE_MASTER_URL = f"{BASE_URL}/public/MatterManagement/CaseMaster.aspx"
VIEW_DOC_URL = f"{BASE_URL}/public/Common/ViewDoc.aspx"
DEFAULT_OUTPUT = "ELG_Database.xlsx"
PAGE_LOAD_TIMEOUT = 60
GRID_LOAD_TIMEOUT = 45

ALLOWED_DOC_TYPES = [
    "briefs",
    "joint proposals and stipulations",
    "orders",
    "petitions",
    "recommended decisions",
    "testimony",
    "transcripts",
]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def parse_case_input(raw: str) -> str:
    raw = raw.strip().rstrip("_")
    if raw.startswith("http"):
        return raw.split("__")[0]
    return f"{CASE_MASTER_URL}?MatterCaseNo={raw}"


def load_cases_from_file(filepath: str) -> list:
    """Read case numbers/URLs from a text file, one per line."""
    if not os.path.exists(filepath):
        print(f"[ERROR] File not found: {filepath}")
        sys.exit(1)
    cases = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#"):
                cases.append(line)
    return cases


def build_driver(headless: bool = True) -> webdriver.Chrome:
    opts = Options()
    if headless:
        opts.add_argument("--headless=new")
    opts.add_argument("--no-sandbox")
    opts.add_argument("--disable-dev-shm-usage")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--window-size=1920,1200")
    opts.add_argument("--disable-blink-features=AutomationControlled")
    opts.add_argument("--log-level=3")
    opts.add_argument("--silent")
    opts.add_experimental_option("excludeSwitches", ["enable-automation", "enable-logging"])
    opts.add_experimental_option("useAutomationExtension", False)
    opts.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
    )

    driver = webdriver.Chrome(options=opts)

    driver.execute_cdp_cmd(
        "Page.addScriptToEvaluateOnNewDocument",
        {"source": "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"},
    )
    return driver


def wait_for_ajax(driver, timeout=30):
    try:
        WebDriverWait(driver, timeout).until(
            lambda d: d.execute_script(
                "return (typeof Sys === 'undefined' || "
                "Sys.WebForms.PageRequestManager.getInstance().get_isInAsyncPostBack() === false)"
            )
        )
    except Exception:
        pass
    time.sleep(2)


def safe_text(element) -> str:
    try:
        return element.text.strip()
    except StaleElementReferenceException:
        return ""


def is_allowed_doc_type(doc_type: str) -> bool:
    lower = doc_type.lower().strip()
    for allowed in ALLOWED_DOC_TYPES:
        if allowed in lower or lower in allowed:
            return True
    return False


# ---------------------------------------------------------------------------
# Case info extraction
# ---------------------------------------------------------------------------

def extract_case_info(driver) -> dict:
    info = {
        "company_name": "",
        "company_type": "",
        "case_number": "",
    }

    wait = WebDriverWait(driver, PAGE_LOAD_TIMEOUT)

    # Wait for the case info panel to load
    try:
        wait.until(EC.presence_of_element_located(
            (By.ID, "GridPlaceHolder_MatterControl1_lblCaseNumberValue")
        ))
    except TimeoutException:
        print("        [ERROR] Timed out waiting for case info to load.")
        # Fallback: try URL params
        url = driver.current_url
        for param in ["MNO=", "MatterCaseNo="]:
            m = re.search(param + r'([^&]+)', url)
            if m:
                info["case_number"] = m.group(1)
                break
        return info

    time.sleep(2)

    # Case Number
    try:
        el = driver.find_element(By.ID, "GridPlaceHolder_MatterControl1_lblCaseNumberValue")
        info["case_number"] = safe_text(el)
    except NoSuchElementException:
        pass

    # Industry Affected (Company Type)
    try:
        el = driver.find_element(By.ID, "GridPlaceHolder_MatterControl1_lblIndustryAffectedValue")
        info["company_type"] = safe_text(el)
    except NoSuchElementException:
        pass

    # Company/Organization — names are inside <li> tags
    try:
        el = driver.find_element(By.ID, "GridPlaceHolder_MatterControl1_pnlCompanyValue")
        items = el.find_elements(By.TAG_NAME, "li")
        names = [safe_text(li) for li in items if safe_text(li)]
        info["company_name"] = "; ".join(names)
    except NoSuchElementException:
        pass

    # Fallback for case number from URL
    if not info["case_number"]:
        url = driver.current_url
        for param in ["MNO=", "MatterCaseNo="]:
            m = re.search(param + r'([^&]+)', url)
            if m:
                info["case_number"] = m.group(1)
                break

    # Fallback for company type from case number
    if not info["company_type"] and info["case_number"]:
        cn = info["case_number"]
        if "-G-" in cn and "-E-" in cn:
            info["company_type"] = "Electric and Gas"
        elif "-G-" in cn:
            info["company_type"] = "Gas"
        elif "-E-" in cn:
            info["company_type"] = "Electric"

    return info


# ---------------------------------------------------------------------------
# Document grid scraping
# ---------------------------------------------------------------------------

def click_filed_documents_tab(driver):
    attempts = [
        "//a[contains(text(),'Filed Documents')]",
        "//a[contains(@id,'lbtPubDoc')]",
        "//*[contains(@id,'PubDoc')]",
    ]
    for xpath in attempts:
        try:
            tab = WebDriverWait(driver, 10).until(
                EC.element_to_be_clickable((By.XPATH, xpath))
            )
            driver.execute_script("arguments[0].click();", tab)
            print("        Clicked 'Filed Documents' tab, waiting for grid...")
            time.sleep(5)
            wait_for_ajax(driver, timeout=30)
            return True
        except (TimeoutException, NoSuchElementException):
            continue

    try:
        driver.execute_script("__doPostBack('ctl00$GridPlaceHolder$lbtPubDoc','')")
        print("        Triggered Filed Documents postback via JS, waiting...")
        time.sleep(5)
        wait_for_ajax(driver, timeout=30)
        return True
    except Exception:
        pass

    print("        [NOTE] Could not find Filed Documents tab — grid may already be visible.")
    return False


def find_document_table(driver):
    id_patterns = [
        "grdFiledDoc_DgdCustomGrid",
        "grdFiledDoc",
        "DgdCustomGrid",
        "GridView",
    ]
    for pattern in id_patterns:
        try:
            tables = driver.find_elements(By.CSS_SELECTOR, f"table[id*='{pattern}']")
            for t in tables:
                rows = t.find_elements(By.TAG_NAME, "tr")
                if len(rows) >= 2:
                    return t
        except Exception:
            continue

    all_tables = driver.find_elements(By.TAG_NAME, "table")
    for t in all_tables:
        try:
            header_text = t.text[:500]
            if "Date Filed" in header_text and "Document Title" in header_text:
                return t
        except StaleElementReferenceException:
            continue

    page_src = driver.page_source
    if "Date Filed" in page_src:
        iframes = driver.find_elements(By.TAG_NAME, "iframe")
        for iframe in iframes:
            try:
                driver.switch_to.frame(iframe)
                tables = driver.find_elements(By.TAG_NAME, "table")
                for t in tables:
                    if "Date Filed" in t.text[:500]:
                        return t
                driver.switch_to.default_content()
            except Exception:
                driver.switch_to.default_content()

    return None


def scrape_current_page(driver, table) -> list:
    rows = []
    all_rows = table.find_elements(By.TAG_NAME, "tr")

    for tr in all_rows:
        cells = tr.find_elements(By.TAG_NAME, "td")
        if len(cells) < 6:
            continue

        first_text = safe_text(cells[0])
        if first_text.lower() in ("", "filters:", "sr.no.", "sr.no"):
            continue
        if not first_text:
            continue
        if not first_text[0].isdigit():
            row_text = safe_text(tr)
            if any(kw in row_text for kw in ["Page Size", "First", "Prev", "Next", "Last"]):
                continue
            continue

        try:
            date_filed = safe_text(cells[1])
            doc_type = safe_text(cells[2])
            doc_title = safe_text(cells[3])
            filing_on_behalf = safe_text(cells[4])
            file_name = safe_text(cells[6]) if len(cells) > 6 else ""

            file_size = ""
            for idx in range(len(cells) - 1, 5, -1):
                txt = safe_text(cells[idx])
                if txt and ("KB" in txt or "MB" in txt or "GB" in txt or "bytes" in txt.lower()):
                    file_size = txt
                    break
            if not file_size and len(cells) > 8:
                file_size = safe_text(cells[8])

            if not is_allowed_doc_type(doc_type):
                continue

            doc_link = ""
            for cell_idx in [3, 6]:
                if cell_idx >= len(cells):
                    continue
                links = cells[cell_idx].find_elements(By.TAG_NAME, "a")
                for a in links:
                    href = a.get_attribute("href") or ""
                    if "ViewDoc.aspx" in href or "DocRefId" in href:
                        doc_link = href
                        break
                    if "MatterFilingItem" in href:
                        doc_link = href
                        break
                if doc_link:
                    break

            if not doc_link:
                for cell in cells:
                    links = cell.find_elements(By.TAG_NAME, "a")
                    for a in links:
                        href = a.get_attribute("href") or ""
                        if "ViewDoc" in href or "DocRefId" in href or "FilingItem" in href:
                            doc_link = href
                            break
                    if doc_link:
                        break

            if not doc_link:
                row_html = tr.get_attribute("innerHTML")
                ref_match = re.search(r'DocRefId=([^&"\'<>\s]+)', row_html)
                if ref_match:
                    doc_id = ref_match.group(1)
                    if not doc_id.startswith("{"):
                        doc_id = "{" + doc_id + "}"
                    doc_link = f"{VIEW_DOC_URL}?DocRefId={doc_id}"
                else:
                    filing_match = re.search(
                        r'FilingSeq=(\d+).*?MatterSeq=(\d+)', row_html
                    )
                    if filing_match:
                        doc_link = (
                            f"{BASE_URL}/public/MatterManagement/MatterFilingItem.aspx"
                            f"?FilingSeq={filing_match.group(1)}"
                            f"&MatterSeq={filing_match.group(2)}"
                        )

            rows.append({
                "date_filed": date_filed,
                "doc_type": doc_type,
                "doc_title": doc_title,
                "filing_on_behalf": filing_on_behalf,
                "file_name": file_name,
                "file_size": file_size,
                "link": doc_link,
            })

        except (IndexError, StaleElementReferenceException):
            continue

    return rows


def get_total_pages(driver) -> int:
    try:
        page_source = driver.page_source
        m = re.search(r'of\s+(\d+)\s', page_source)
        if m:
            return int(m.group(1))
        pager_links = driver.find_elements(By.XPATH, "//a[contains(@href,'Page$')]")
        if pager_links:
            nums = []
            for pl in pager_links:
                txt = safe_text(pl)
                if txt.isdigit():
                    nums.append(int(txt))
            if nums:
                return max(nums)
    except Exception:
        pass
    return 1


def click_next_page(driver) -> bool:
    try:
        next_links = driver.find_elements(By.XPATH, "//a[contains(text(),'Next')]")
        for nl in next_links:
            try:
                parent_row = nl.find_element(By.XPATH, "./ancestor::tr[1]")
                row_text = safe_text(parent_row)
                if "Page Size" in row_text or "First" in row_text or "Prev" in row_text:
                    driver.execute_script("arguments[0].click();", nl)
                    time.sleep(3)
                    wait_for_ajax(driver, timeout=20)
                    return True
            except Exception:
                continue
        return False
    except Exception:
        return False


def set_page_size_max(driver):
    try:
        selects = driver.find_elements(
            By.XPATH,
            "//select[contains(@id,'PageSize') or contains(@id,'pageSize') or contains(@id,'ddlPageSize')]"
        )
        if selects:
            from selenium.webdriver.support.ui import Select
            sel = Select(selects[0])
            options = [o.text.strip() for o in sel.options]
            nums = [int(o) for o in options if o.isdigit()]
            if nums:
                largest = str(max(nums))
                print(f"        Setting page size to {largest}...")
                sel.select_by_visible_text(largest)
                time.sleep(5)
                wait_for_ajax(driver, timeout=30)
    except Exception:
        pass


def scrape_all_documents(driver) -> list:
    all_docs = []

    table = find_document_table(driver)
    if table is None:
        return all_docs

    set_page_size_max(driver)

    table = find_document_table(driver)
    if table is None:
        return all_docs

    page_num = 1
    total = get_total_pages(driver)
    print(f"  Detected up to {total} page(s) of documents.")

    while True:
        print(f"  Scraping page {page_num}...", end=" ")

        table = find_document_table(driver)
        if table is None:
            print("table not found after page change.")
            break

        docs = scrape_current_page(driver, table)
        print(f"found {len(docs)} matching documents.")
        all_docs.extend(docs)

        if page_num >= total:
            break
        if not click_next_page(driver):
            break
        page_num += 1

    return all_docs


# ---------------------------------------------------------------------------
# Excel writing
# ---------------------------------------------------------------------------

def append_to_excel(output_path: str, case_info: dict, documents: list):
    if os.path.exists(output_path):
        wb = load_workbook(output_path)
        ws = wb.active
    else:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        headers = [
            "Case Number", "Company Name", "Company Type", "Date Filed",
            "Document Type", "Document Title", "Filing on Behalf Of",
            "File Name", "File Size", "Link to Document",
        ]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
            ws.cell(row=1, column=col).font = Font(bold=True)

    next_row = ws.max_row + 1
    if next_row == 2 and ws.cell(row=2, column=1).value is None:
        next_row = 2

    link_font = Font(color="0563C1", underline="single")

    for doc in documents:
        row = next_row
        ws.cell(row=row, column=1, value=case_info["case_number"])
        ws.cell(row=row, column=2, value=case_info["company_name"])
        ws.cell(row=row, column=3, value=case_info["company_type"])
        ws.cell(row=row, column=4, value=doc["date_filed"])
        ws.cell(row=row, column=5, value=doc["doc_type"])
        ws.cell(row=row, column=6, value=doc["doc_title"])
        ws.cell(row=row, column=7, value=doc["filing_on_behalf"])
        ws.cell(row=row, column=8, value=doc["file_name"])
        ws.cell(row=row, column=9, value=doc["file_size"])

        link = doc.get("link", "")
        if link:
            ws.cell(row=row, column=10).hyperlink = link
            ws.cell(row=row, column=10).value = link
            ws.cell(row=row, column=10).font = link_font
        else:
            ws.cell(row=row, column=10, value="")

        next_row += 1

    wb.save(output_path)
    return len(documents)


# ---------------------------------------------------------------------------
# Single case processor
# ---------------------------------------------------------------------------

def process_single_case(driver, case_input: str, output_path: str) -> dict:
    """Process one case. Returns a result dict with status and counts."""
    url = parse_case_input(case_input)
    result = {"input": case_input, "url": url, "case_number": "", "docs": 0, "status": ""}

    print(f"  Loading: {url}")
    driver.get(url)

    print("        Waiting for page to render...")
    time.sleep(10)
    wait_for_ajax(driver, timeout=30)

    page_text = driver.page_source
    if "does not exist" in page_text:
        print("        [FAILED] Case number does not exist.")
        result["status"] = "NOT FOUND"
        return result

    print("        Extracting case info...")
    case_info = extract_case_info(driver)
    result["case_number"] = case_info["case_number"]

    print(f"        Company:  {case_info['company_name']}")
    print(f"        Type:     {case_info['company_type']}")
    print(f"        Case No:  {case_info['case_number']}")

    print("        Opening Filed Documents...")
    click_filed_documents_tab(driver)

    print("        Waiting for document grid...")
    time.sleep(10)
    wait_for_ajax(driver, timeout=30)

    table = find_document_table(driver)
    if table is None:
        print("        [FAILED] Document grid did not load.")
        screenshot_path = f"debug_{case_info['case_number'] or 'unknown'}.png"
        driver.save_screenshot(screenshot_path)
        print(f"        Screenshot saved to: {screenshot_path}")
        result["status"] = "GRID FAILED"
        return result

    documents = scrape_all_documents(driver)

    if not documents:
        print("        [NOTE] No documents matched the allowed types.")
        result["status"] = "NO MATCHES"
        return result

    count = append_to_excel(output_path, case_info, documents)
    result["docs"] = count
    result["status"] = "OK"
    print(f"        [OK] {count} documents saved.")
    return result


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Scrape NY DPS rate case documents into an Excel database.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  Single case:
    python dps_scraper.py 25-G-0073

  Multiple cases:
    python dps_scraper.py 25-G-0073 23-G-0225 24-E-0165

  From a text file (one case per line):
    python dps_scraper.py --file cases.txt

  Mix and match:
    python dps_scraper.py 25-G-0073 "https://documents.dps.ny.gov/..." --output ./ELG_Database.xlsx

  With visible browser:
    python dps_scraper.py 25-G-0073 23-G-0225 --visible
        """,
    )
    parser.add_argument(
        "cases", nargs="*", default=[],
        help="One or more DPS case URLs or case numbers (e.g. 25-G-0073)",
    )
    parser.add_argument(
        "--file", "-f", default=None,
        help="Path to a text file with one case number/URL per line",
    )
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT, help="Path to output Excel file")
    parser.add_argument("--visible", action="store_true", help="Show the browser window for debugging")
    args = parser.parse_args()

    # Collect all cases from args + file
    all_cases = list(args.cases)
    if args.file:
        file_cases = load_cases_from_file(args.file)
        all_cases.extend(file_cases)

    if not all_cases:
        parser.print_help()
        print("\n[ERROR] No cases provided. Pass case numbers as arguments or use --file.")
        sys.exit(1)

    output_path = args.output

    print("=" * 60)
    print("  NY DPS Rate Case Document Scraper")
    print("=" * 60)
    print(f"  Cases:  {len(all_cases)}")
    print(f"  Output: {output_path}")
    print()
    print("  Filtering for document types:")
    for dt in ALLOWED_DOC_TYPES:
        print(f"    - {dt.title()}")
    print()
    print("  Cases to process:")
    for c in all_cases:
        print(f"    - {c}")
    print()

    print("[SETUP] Launching browser...")
    driver = build_driver(headless=not args.visible)
    driver.set_page_load_timeout(PAGE_LOAD_TIMEOUT)

    results = []

    try:
        for i, case_input in enumerate(all_cases, 1):
            print()
            print("-" * 60)
            print(f"  CASE {i}/{len(all_cases)}: {case_input}")
            print("-" * 60)

            try:
                result = process_single_case(driver, case_input, output_path)
                results.append(result)
            except Exception as e:
                print(f"        [ERROR] {type(e).__name__}: {e}")
                results.append({
                    "input": case_input,
                    "url": "",
                    "case_number": "",
                    "docs": 0,
                    "status": f"ERROR: {e}",
                })

            # Brief pause between cases
            if i < len(all_cases):
                print("        Pausing before next case...")
                time.sleep(3)

    finally:
        driver.quit()

    # ---- Final summary ----
    print()
    print()
    print("=" * 60)
    print("  FINAL SUMMARY")
    print("=" * 60)
    print()

    total_docs = 0
    ok_count = 0
    fail_count = 0

    for r in results:
        status_icon = "OK" if r["status"] == "OK" else "!!"
        case_label = r["case_number"] or r["input"]
        print(f"  [{status_icon}] {case_label:20s}  {r['docs']:4d} docs   ({r['status']})")
        total_docs += r["docs"]
        if r["status"] == "OK":
            ok_count += 1
        else:
            fail_count += 1

    print()
    print(f"  Total cases:     {len(results)}")
    print(f"  Succeeded:       {ok_count}")
    print(f"  Failed/skipped:  {fail_count}")
    print(f"  Total documents: {total_docs}")
    print(f"  Output file:     {output_path}")
    print()


if __name__ == "__main__":
    main()
