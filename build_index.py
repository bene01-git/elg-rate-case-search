#!/usr/bin/env python3
from __future__ import annotations
"""
Build Search Index
===================
Reads text files and ELG_Database.xlsx, outputs a compact search_data.json
for the static GitHub Pages search site.

Usage:
    python build_index.py
    python build_index.py --text-dir searchable_pdfs/text --output site/search_data.json

This is run by GitHub Actions on every push, or manually before deploying.
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


DEFAULT_TEXT_DIR = os.path.join("searchable_pdfs", "text")
DEFAULT_EXCEL = "ELG_Database.xlsx"
DEFAULT_OUTPUT = os.path.join("site", "search_data.json")


def load_metadata(excel_path):
    if not HAS_OPENPYXL or not os.path.exists(excel_path):
        return []

    wb = load_workbook(excel_path)
    ws = wb.active
    records = []

    for row_idx in range(2, ws.max_row + 1):
        case_number = str(ws.cell(row=row_idx, column=1).value or "").strip()
        company_name = str(ws.cell(row=row_idx, column=2).value or "").strip()
        company_type = str(ws.cell(row=row_idx, column=3).value or "").strip()
        date_filed = str(ws.cell(row=row_idx, column=4).value or "").strip()
        doc_type = str(ws.cell(row=row_idx, column=5).value or "").strip()
        doc_title = str(ws.cell(row=row_idx, column=6).value or "").strip()
        filing_on_behalf = str(ws.cell(row=row_idx, column=7).value or "").strip()

        cell = ws.cell(row=row_idx, column=10)
        link = ""
        if cell.hyperlink:
            link = cell.hyperlink.target or ""
        if not link:
            link = str(cell.value or "")

        item_no = ""
        m = re.search(r'FilingSeq=(\d+)', link)
        if m:
            item_no = m.group(1)

        if not case_number and not doc_title:
            continue

        records.append({
            "case_number": case_number,
            "company_name": company_name,
            "company_type": company_type,
            "date_filed": date_filed,
            "doc_type": doc_type,
            "doc_title": doc_title,
            "filing_on_behalf": filing_on_behalf,
            "link": link,
            "item_no": item_no,
        })

    wb.close()
    return records


def match_file_to_metadata(filename, records):
    name = filename.replace(".txt", "").replace(".pdf", "").lower()
    best_match = None
    best_score = 0
    for rec in records:
        score = 0
        if rec["case_number"] and rec["case_number"].lower() in name:
            score += 3
        if rec["doc_title"]:
            t = rec["doc_title"][:60].lower()
            if t in name:
                score += 5
            elif t[:30] in name:
                score += 3
        if rec["doc_type"] and rec["doc_type"].lower() in name:
            score += 2
        if rec["company_name"]:
            c = rec["company_name"].split(",")[0].lower()
            if c in name:
                score += 1
        if score > best_score:
            best_score = score
            best_match = rec
    return best_match if best_score >= 5 else None


def format_citation(meta):
    parts = []
    if meta.get("case_number"):
        parts.append(f"Case {meta['case_number']}")
    if meta.get("doc_type") and meta.get("doc_title"):
        parts.append(f"{meta['doc_type']}: {meta['doc_title']}")
    elif meta.get("doc_title"):
        parts.append(meta["doc_title"])
    if meta.get("filing_on_behalf"):
        parts.append(f"filed by {meta['filing_on_behalf']}")
    if meta.get("company_name"):
        if meta.get("filing_on_behalf", "").lower() != meta["company_name"].lower():
            parts.append(f"on behalf of {meta['company_name']}")
    if meta.get("item_no"):
        parts.append(f"Paper No. {meta['item_no']}")
    if meta.get("date_filed"):
        parts.append(f"({meta['date_filed']})")
    return ", ".join(parts)


def split_sentences(text):
    text = re.sub(r'---\s*Page\s+\d+\s*---', ' ', text)
    text = re.sub(r'\n+', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    sentences = re.split(r'(?<=[.!?])\s+(?=[A-Z"])', text)
    result = []
    for s in sentences:
        s = s.strip()
        if not s or len(s) < 10:
            continue
        if len(s) > 800:
            parts = re.split(r'[;:]\s+', s)
            result.extend(p.strip() for p in parts if p.strip() and len(p.strip()) >= 10)
        else:
            result.append(s)
    return result


def main():
    parser = argparse.ArgumentParser(description="Build search index JSON.")
    parser.add_argument("--text-dir", default=DEFAULT_TEXT_DIR)
    parser.add_argument("--excel", default=DEFAULT_EXCEL)
    parser.add_argument("--output", "-o", default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if not os.path.isdir(args.text_dir):
        print(f"[ERROR] Text directory not found: {args.text_dir}")
        sys.exit(1)

    print(f"Loading metadata from {args.excel}...")
    records = load_metadata(args.excel)
    print(f"Loaded {len(records)} metadata records.")

    txt_files = sorted(Path(args.text_dir).glob("*.txt"))
    print(f"Found {len(txt_files)} text files.")

    documents = []

    for i, filepath in enumerate(txt_files):
        try:
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        except Exception:
            continue

        if not content.strip():
            continue

        sentences = split_sentences(content)
        if not sentences:
            continue

        meta = match_file_to_metadata(filepath.name, records)
        citation = format_citation(meta) if meta else ""

        doc = {
            "id": i,
            "filename": filepath.name,
            "case_number": meta.get("case_number", "") if meta else "",
            "company_name": meta.get("company_name", "") if meta else "",
            "doc_type": meta.get("doc_type", "") if meta else "",
            "doc_title": meta.get("doc_title", "") if meta else "",
            "filing_on_behalf": meta.get("filing_on_behalf", "") if meta else "",
            "date_filed": meta.get("date_filed", "") if meta else "",
            "item_no": meta.get("item_no", "") if meta else "",
            "link": meta.get("link", "") if meta else "",
            "citation": citation,
            "sentences": sentences,
        }
        documents.append(doc)

        if (i + 1) % 50 == 0:
            print(f"  Processed {i + 1}/{len(txt_files)} files...")

    os.makedirs(os.path.dirname(args.output), exist_ok=True)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump({"documents": documents, "count": len(documents)}, f)

    size_mb = os.path.getsize(args.output) / (1024 * 1024)
    print(f"\nDone: {len(documents)} documents → {args.output} ({size_mb:.1f} MB)")


if __name__ == "__main__":
    main()
